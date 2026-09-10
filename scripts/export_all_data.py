import os
import re
import json
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

XLSX_PATH = 'Skyrim True Believer.xlsx'
DATA_DIR = 'data'
os.makedirs(DATA_DIR, exist_ok=True)

print("Loading workbook for sheet extraction...")
wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)

# -------------------------------------------------------------
# 1. NAVIGATION
# -------------------------------------------------------------
print("Exporting navigation...")
navigation_data = {
    "title": "Skyrim True Believer",
    "subtitle": "Полная интерактивная база данных и энциклопедия сборки STB 3.0",
    "version": "3.0",
    "update_note": "Изменения текущей версии во вкладках окрашены зеленым",
    "test_room": "coc dev",
    "developers": [
        {"name": "Sneyk", "role": "Основной разработчик"},
        {"name": "MoonDream", "role": "Основной разработчик"},
        {"name": "Polaris", "role": "Основной разработчик"},
        {"name": "KBA3AP", "role": "Основной разработчик"},
        {"name": "Мельче", "role": "Разработчик"},
        {"name": "Nikita", "role": "Разработчик"},
        {"name": "Frem", "role": "Разработчик"}
    ],
    "helpers": ["StarkMP", "Мастер Изгой Егор", "Sewhass"],
    "artists": ["SUNBARDO", "silk_911", "leonhetch"],
    "sponsors": ["Глубже", "silk_911", "leonhetch"],
    "socials": [
        {"name": "Boosty", "url": "https://boosty.to/skyrimtruebeliever", "icon": "assets/img/image1.png"},
        {"name": "Discord", "url": "https://discord.gg/", "icon": "assets/img/image2.png"},
        {"name": "Patreon", "url": "https://www.patreon.com/", "icon": "assets/img/image3.png"}
    ],
    "banner": "assets/img/image30.png"
}
with open(f"{DATA_DIR}/navigation.json", "w", encoding="utf-8") as f:
    json.dump(navigation_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 2. BASE MECHANICS
# -------------------------------------------------------------
print("Exporting base mechanics...")
ws_base = wb['📚База']
base_sections = []
intoxication_table = []

for r, row in enumerate(ws_base.iter_rows(values_only=True)):
    for c in [12, 47]:
        if c < len(row):
            val = str(row[c] or '').strip()
            if val and len(val) > 10:
                lines = [l.strip() for l in val.split('\n') if l.strip()]
                if lines:
                    title = lines[0]
                    content = '\n'.join(lines[1:])
                    # avoid duplicates
                    if not any(s['title'] == title for s in base_sections):
                        base_sections.append({
                            "id": f"sec-{len(base_sections)+1}",
                            "title": title,
                            "content": content
                        })

# Intoxication penalty table
intoxication_table = [
    {"level": "1 уровень", "range": "60–69", "penalty": "-20% эффективности зелий", "recommendation": "Безопасная зона для боя"},
    {"level": "2 уровень", "range": "70–79", "penalty": "-40% эффективности зелий", "recommendation": "Ощутимое снижение эффекта"},
    {"level": "3 уровень", "range": "80–89", "penalty": "-60% эффективности зелий", "recommendation": "Критическое состояние"},
    {"level": "4 уровень", "range": "90–99", "penalty": "-80% эффективности зелий", "recommendation": "Зелья практически бесполезны"},
    {"level": "5 уровень", "range": "100+", "penalty": "-100% эффективности зелий", "recommendation": "Полная невосприимчивость к зельям"}
]

base_data = {
    "sections": base_sections,
    "intoxication_table": intoxication_table,
    "dev_room": {
        "command": "coc dev",
        "description": "Позволяет попасть в тестовую комнату со стендами, алтарями и манекенами для тестирования 99% контента сборки."
    }
}
with open(f"{DATA_DIR}/base.json", "w", encoding="utf-8") as f:
    json.dump(base_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 3. RACES & STANDING STONES
# -------------------------------------------------------------
print("Exporting races & standing stones...")
ws_races = wb['🗿РасыКамни']
races_stones_data = {
    "start_rules": [
        "Выбор камня судьбы на старте",
        "Распределение 250 атрибутов (не более 200 в один атрибут) и 250% регенераций",
        "Выбор расходуемого атрибута (запас сил или магия), который тратится на атаки оружием, уворот, прыжок, спринт и плавание"
    ],
    "innate_abilities": [],
    "standing_stones": []
}

# Innate abilities from R9 to R27
for r, row in enumerate(ws_races.iter_rows(values_only=True)):
    row_num = r + 1
    if 9 <= row_num <= 27:
        for c in [14, 21, 28, 35]:
            if c < len(row):
                val = str(row[c] or '').strip()
                if val:
                    races_stones_data["innate_abilities"].append(val)

# Standing stones from R33, R40, R47, R54
stone_rows = [
    (33, 34, "Группа Воина"),
    (40, 41, "Группа Вора"),
    (47, 48, "Группа Мага"),
    (54, 55, "Особые камни")
]
rows_list = list(ws_races.iter_rows(values_only=True))
for name_r, desc_r, group_name in stone_rows:
    row_names = rows_list[name_r - 1]
    row_descs = rows_list[desc_r - 1]
    for c in [3, 13, 23, 33, 43]:
        if c < len(row_names):
            name = str(row_names[c] or '').strip()
            desc = str(row_descs[c] or '').strip() if c < len(row_descs) else ''
            if name and name != 'None':
                races_stones_data["standing_stones"].append({
                    "name": name,
                    "group": group_name,
                    "description": desc
                })

with open(f"{DATA_DIR}/races_stones.json", "w", encoding="utf-8") as f:
    json.dump(races_stones_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 4. AEDRA
# -------------------------------------------------------------
print("Exporting Aedra...")
ws_aedra = wb['😇Аэдра']
aedra_images = {
    "Аркей": "assets/img/image7.png",
    "Джулианос": "assets/img/image10.png",
    "Кинарет": "assets/img/image15.png",
    "Акатош": "assets/img/image9.png",
    "Стендарр": "assets/img/image28.png",
    "Талос": "assets/img/image14.png",
    "Зенитар": "assets/img/image12.png",
    "Дибелла": "assets/img/image13.png",
    "Мара": "assets/img/image16.png"
}
aedra_list = []
aedra_rules = [
    "Можно поклоняться трем Аэдра одновременно",
    "Принятие Аэдра открывается на 5, 25 и 45 уровне персонажа",
    "Молитвы совершаются на священных алтарях",
    "Божественный дар и Благословение усиливаются за каждую совершенную молитву (до 5 молитв на бога)",
    "Принятие Аэдра сбрасывает прогресс поклонений Даэдра",
    "Карма дополнительно повышает силу благословения на 10% за каждые 10 ед. кармы (макс. +100%)",
    "Общая стоимость всех молитв: 1 200 000 золота"
]

prayer_costs = [
    {"aedra_slot": "1-й Аэдра", "prayers": [
        {"num": 1, "cost": "10 000", "level": 5},
        {"num": 2, "cost": "20 000", "level": 8},
        {"num": 3, "cost": "30 000", "level": 11},
        {"num": 4, "cost": "40 000", "level": 14},
        {"num": 5, "cost": "50 000", "level": 17}
    ]},
    {"aedra_slot": "2-й Аэдра", "prayers": [
        {"num": 1, "cost": "60 000", "level": 25},
        {"num": 2, "cost": "70 000", "level": 28},
        {"num": 3, "cost": "80 000", "level": 31},
        {"num": 4, "cost": "90 000", "level": 34},
        {"num": 5, "cost": "100 000", "level": 37}
    ]},
    {"aedra_slot": "3-й Аэдра", "prayers": [
        {"num": 1, "cost": "110 000", "level": 45},
        {"num": 2, "cost": "120 000", "level": 48},
        {"num": 3, "cost": "130 000", "level": 51},
        {"num": 4, "cost": "140 000", "level": 54},
        {"num": 5, "cost": "150 000", "level": 57}
    ]}
]

for r, row in enumerate(ws_aedra.iter_rows(values_only=True)):
    name = str(row[3] or '').strip() if len(row) > 3 else ''
    if name in aedra_images:
        gift = str(row[4] or '').strip() if len(row) > 4 else ''
        blessing = str(row[5] or '').strip() if len(row) > 5 else ''
        amulet = str(row[6] or '').strip() if len(row) > 6 else ''
        backpack = str(row[7] or '').strip() if len(row) > 7 else ''
        notes = str(row[9] or '').strip() if len(row) > 9 else ''
        
        aedra_list.append({
            "name": name,
            "image": aedra_images.get(name, ""),
            "gift": gift,
            "blessing": blessing,
            "amulet": amulet,
            "backpack": backpack,
            "notes": notes
        })

with open(f"{DATA_DIR}/aedra.json", "w", encoding="utf-8") as f:
    json.dump({"rules": aedra_rules, "prayer_costs": prayer_costs, "gods": aedra_list}, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 5. DAEDRA
# -------------------------------------------------------------
print("Exporting Daedra...")
ws_daedra = wb['😈Даэдра']
daedra_images = {
    "Азура": "assets/img/image23.jpg",
    "Боэтия": "assets/img/image11.jpg",
    "Вермина": "assets/img/image6.jpg",
    "Клавикус Вайл": "assets/img/image8.jpg",
    "Меридия": "assets/img/image20.jpg",
    "Молаг Бал": "assets/img/image25.jpg",
    "Мефала": "assets/img/image24.jpg",
    "Мерунес Дагон": "assets/img/image18.jpg",
    "Малакат": "assets/img/image21.jpg",
    "Намира": "assets/img/image17.jpg",
    "Ноктюрнал": "assets/img/image26.jpg",
    "Периайт": "assets/img/image22.jpg",
    "Сангвин": "assets/img/image27.jpg",
    "Хирсин": "assets/img/image32.png",
    "Шеогорат": "assets/img/image29.png",
    "Хермеус Мора": "assets/img/image37.png"
}
daedra_rules = [
    "Можно поклоняться трем Даэдра одновременно",
    "Принятие Даэдра открывается на 5, 25 и 45 уровне",
    "Подношения совершаются на алтарях или особым талантом поклонения",
    "Божественный дар усиливается за каждое подношение",
    "Локации святилищ открываются при прочтении соответствующих книг или нахождении алтарей",
    "Возврат артефакта дает мощный дополнительный бонус (4-5 зачарований)",
    "Общая стоимость подношений: 1 200 000 золота"
]

daedra_list = []
current_daedra = None
for r, row in enumerate(ws_daedra.iter_rows(values_only=True)):
    name_cand = str(row[3] or '').strip() if len(row) > 3 else ''
    if name_cand in daedra_images:
        gift = str(row[4] or '').strip() if len(row) > 4 else ''
        art_name = str(row[5] or '').strip() if len(row) > 5 else ''
        art_effect = str(row[7] or '').strip() if len(row) > 7 else ''
        damage_armor = str(row[8] or '').strip() if len(row) > 8 else ''
        weight = str(row[9] or '').strip() if len(row) > 9 else ''
        quest = str(row[10] or '').strip() if len(row) > 10 else ''
        sphere = str(row[12] or '').strip() if len(row) > 12 else ''
        
        current_daedra = {
            "name": name_cand,
            "image": daedra_images.get(name_cand, ""),
            "gift": gift,
            "artifacts": [],
            "quest": quest,
            "sphere": sphere
        }
        if art_name or art_effect:
            current_daedra["artifacts"].append({
                "name": art_name,
                "effect": art_effect,
                "stats": damage_armor,
                "weight": weight
            })
        daedra_list.append(current_daedra)
    elif current_daedra:
        # sub-row for alternative artifact / extra stats
        sub_name = str(row[5] or '').strip() if len(row) > 5 else ''
        sub_effect = str(row[7] or '').strip() if len(row) > 7 else ''
        sub_stats = str(row[8] or '').strip() if len(row) > 8 else ''
        sub_weight = str(row[9] or '').strip() if len(row) > 9 else ''
        if sub_name or sub_effect:
            current_daedra["artifacts"].append({
                "name": sub_name,
                "effect": sub_effect,
                "stats": sub_stats,
                "weight": sub_weight
            })

with open(f"{DATA_DIR}/daedra.json", "w", encoding="utf-8") as f:
    json.dump({"rules": daedra_rules, "gods": daedra_list}, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 6. CURSED (Vampire & Werewolf)
# -------------------------------------------------------------
print("Exporting cursed (vampires & werewolves)...")
with zipfile.ZipFile(XLSX_PATH) as z:
    root6 = ET.fromstring(z.read('xl/drawings/drawing6.xml'))
    ns = {'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    
    shapes = []
    for sp in root6.findall('.//xdr:sp', ns):
        xfrm = sp.find('.//a:xfrm', ns)
        off = xfrm.find('a:off', ns) if xfrm is not None else None
        x = int(off.attrib.get('x', 0)) if off is not None else 0
        y = int(off.attrib.get('y', 0)) if off is not None else 0
        
        tx = sp.find('.//xdr:txBody', ns)
        text = ''
        if tx is not None:
            text = ' '.join([''.join([t.text for t in p.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t') if t.text]) for p in tx.findall('.//a:p', ns)]).strip()
        if text and text != '*':
            shapes.append({"x": x, "y": y, "text": text})

# Separate shapes by x into Werewolf (left / x < 4500000) and Vampire (right / x >= 4500000)
werewolf_shapes = [s for s in shapes if s['x'] < 4500000]
vampire_shapes = [s for s in shapes if s['x'] >= 4500000]

def build_perks_from_shapes(shape_list):
    titles = []
    descs = []
    for s in shape_list:
        t = s['text']
        if 'Всего:' in t or 'без формы' in t or 'Стать' in t:
            continue
        if re.search(r'\(ЧК-[0-9/]+\)', t) or re.search(r'\(ЧК\s*[0-9/]+\)', t):
            titles.append(s)
        else:
            descs.append(s)
    
    perks = []
    for t in titles:
        # find closest description
        best_desc = ""
        min_dist = float('inf')
        for d in descs:
            # description should be close in x and below or near y
            dx = abs(t['x'] - d['x'])
            dy = d['y'] - t['y']
            if -100000 <= dy <= 900000 and dx < 800000:
                dist = dx + dy
                if dist < min_dist:
                    min_dist = dist
                    best_desc = d['text']
        
        m = re.search(r'(.*?)\s*\((ЧК-[0-9/]+)\)', t['text'])
        if m:
            p_name = m.group(1).strip()
            p_req = m.group(2).strip()
        else:
            p_name = t['text']
            p_req = "ЧК-1"
            
        perks.append({
            "name": p_name,
            "req": p_req,
            "effect": best_desc
        })
    return perks

def parse_chk_lvl(req_str):
    nums = re.findall(r'\d+', str(req_str))
    return int(nums[0]) if nums else 0

werewolf_perks_list = build_perks_from_shapes(werewolf_shapes)
werewolf_perks_list.sort(key=lambda p: (parse_chk_lvl(p['req']), p['name']))

vampire_perks_list = build_perks_from_shapes(vampire_shapes)
vampire_perks_list.sort(key=lambda p: (parse_chk_lvl(p['req']), p['name']))

cursed_data = {
    "werewolf": {
        "title": "Вервольф (Оборотень)",
        "rules": [
            "Стать вервольфом можно с 5 уровня персонажа",
            "Навыки требуют показатель чистоты крови (ЧК 1–5), который повышается у даэдрических алтарей и при пожирании сердец",
            "Древо способностей работает в человеческой форме (без обязательного перехода в форму зверя)",
            "Смешение кровей: дает возможность стать Вампиром-гибридом"
        ],
        "perks": werewolf_perks_list
    },
    "vampire": {
        "title": "Вампир (Владыка ночи)",
        "rules": [
            "Стать вампиром можно с 5 уровня персонажа",
            "Навыки требуют показатель чистоты крови (ЧК 1–5), который повышается у даэдрических алтарей и при питье крови",
            "Врожденные способности: Парение, Водное дыхание, Ночное зрение и Чутье Вампира",
            "Штрафы: Уязвимость к солнцу, сопротивление огню -200%, сопротивление холоду +100%",
            "Смешение кровей: дает возможность стать Вервольфом-гибридом"
        ],
        "perks": vampire_perks_list
    }
}
with open(f"{DATA_DIR}/cursed.json", "w", encoding="utf-8") as f:
    json.dump(cursed_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 7. PERK TREES (18 SKILL TREES)
# -------------------------------------------------------------
print("Exporting 18 skill perk trees...")
tree_names_ordered = [
    {"id": "heavy_armor", "name": "Тяжелая броня", "category": "Воин"},
    {"id": "block", "name": "Блокирование", "category": "Воин"},
    {"id": "archery", "name": "Дальний бой (Стрельба)", "category": "Воин"},
    {"id": "sneak", "name": "Скрытность", "category": "Вор"},
    {"id": "voice", "name": "Путь голоса (Ту'ум)", "category": "Особое"},
    {"id": "enchanting", "name": "Зачарование", "category": "Маг"},
    {"id": "alteration", "name": "Изменение", "category": "Маг"},
    {"id": "restoration", "name": "Восстановление", "category": "Маг"},
    {"id": "illusion", "name": "Иллюзия", "category": "Маг"},
    {"id": "speech", "name": "Красноречие", "category": "Вор"},
    {"id": "pickpocket", "name": "Воровство (Карманные кражи)", "category": "Вор"},
    {"id": "destruction", "name": "Разрушение", "category": "Маг"},
    {"id": "alchemy", "name": "Алхимия", "category": "Вор"},
    {"id": "conjuration", "name": "Колдовство", "category": "Маг"},
    {"id": "evasion", "name": "Уклонение (Легкая броня)", "category": "Воин"},
    {"id": "smithing", "name": "Кузнечное дело", "category": "Воин"},
    {"id": "one_handed", "name": "Одноручное оружие", "category": "Воин"},
    {"id": "two_handed", "name": "Двуручное оружие", "category": "Воин"}
]

with zipfile.ZipFile(XLSX_PATH) as z:
    root7 = ET.fromstring(z.read('xl/drawings/drawing7.xml'))
    anchors = root7.findall('xdr:oneCellAnchor', ns)
    
    perks_data = []
    for idx, anc in enumerate(anchors):
        if idx >= len(tree_names_ordered):
            break
        meta = tree_names_ordered[idx]
        grp = anc.find('xdr:grpSp', ns)
        if grp is None:
            continue
        
        shapes = []
        for sp in grp.findall('.//xdr:sp', ns):
            xfrm = sp.find('.//a:xfrm', ns)
            off = xfrm.find('a:off', ns) if xfrm is not None else None
            x = int(off.attrib.get('x', 0)) if off is not None else 0
            y = int(off.attrib.get('y', 0)) if off is not None else 0
            
            tx = sp.find('.//xdr:txBody', ns)
            text = ''
            if tx is not None:
                text = ' '.join([''.join([t.text for t in p.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t') if t.text]) for p in tx.findall('.//a:p', ns)]).strip()
            if text and text != '*' and text != '+-' and text != '**' and text != 'ё':
                shapes.append({"x": x, "y": y, "text": text})
        
        # Partition shapes into title vs description
        titles = []
        descs = []
        for s in shapes:
            t = s['text']
            if 'Всего:' in t or t == meta['name'] or t == 'F5' or 'таблица сломалась' in t:
                continue
            # Titles typically have level requirement in parenthesis e.g. (10/20) or short title
            if re.search(r'\([0-9/]+\)', t) or (len(t) < 35 and not any(k in t for k in ['+', '%', 'урон', 'снижает', 'эффект'])):
                titles.append(s)
            else:
                descs.append(s)
        
        tree_perks = []
        for t in titles:
            best_desc = ""
            min_dist = float('inf')
            for d in descs:
                dx = abs(t['x'] - d['x'])
                dy = d['y'] - t['y']
                if -150000 <= dy <= 950000 and dx < 850000:
                    dist = dx * 1.5 + dy
                    if dist < min_dist:
                        min_dist = dist
                        best_desc = d['text']
            
            # parse req
            m = re.search(r'(.*?)\s*\(([0-9/]+)\)', t['text'])
            if m:
                p_name = m.group(1).strip()
                p_lvl = m.group(2).strip()
            else:
                p_name = t['text']
                p_lvl = "1"
            
            tree_perks.append({
                "name": p_name,
                "level": p_lvl,
                "description": best_desc if best_desc else "Базовый перк ветки"
            })
            
        # Sort perks by required level ascending (1 -> 100)
        def parse_min_lvl(lvl_str):
            nums = re.findall(r'\d+', str(lvl_str))
            return int(nums[0]) if nums else 0
            
        tree_perks.sort(key=lambda p: (parse_min_lvl(p['level']), p['name']))
            
        perks_data.append({
            "id": meta["id"],
            "name": meta["name"],
            "category": meta["category"],
            "total_perks": len(tree_perks),
            "perks": tree_perks
        })

with open(f"{DATA_DIR}/perks.json", "w", encoding="utf-8") as f:
    json.dump(perks_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 8. CONSUMABLES (Расходники)
# -------------------------------------------------------------
print("Exporting consumables...")
ws_cons = wb['🧪Расходники']
potions = []
poisons = []
alcohol = []
food = []
scrolls = []
staves = []

for r, row in enumerate(ws_cons.iter_rows(values_only=True)):
    if r < 3:
        continue
    
    # Potions (col 1..6)
    p_name = str(row[1] or '').strip() if len(row) > 1 else ''
    p_tier = str(row[2] or '').strip() if len(row) > 2 else ''
    p_eff = str(row[3] or '').strip() if len(row) > 3 else ''
    p_dur = str(row[4] or '').strip() if len(row) > 4 else ''
    p_tox = str(row[5] or '').strip() if len(row) > 5 else ''
    p_cost = str(row[6] or '').strip() if len(row) > 6 else ''
    if p_eff:
        potions.append({
            "name": p_name,
            "tier": p_tier,
            "effect": p_eff,
            "duration": p_dur,
            "toxicity": p_tox,
            "cost": p_cost
        })
    
    # Poisons (col 8..12)
    poi_name = str(row[8] or '').strip() if len(row) > 8 else ''
    poi_tier = str(row[9] or '').strip() if len(row) > 9 else ''
    poi_eff = str(row[10] or '').strip() if len(row) > 10 else ''
    poi_dur = str(row[11] or '').strip() if len(row) > 11 else ''
    poi_cost = str(row[12] or '').strip() if len(row) > 12 else ''
    if poi_eff:
        poisons.append({
            "name": poi_name,
            "tier": poi_tier,
            "effect": poi_eff,
            "duration": poi_dur,
            "cost": poi_cost
        })
        
    # Alcohol (col 14..20)
    alc_name = str(row[14] or '').strip() if len(row) > 14 else ''
    alc_eff = str(row[15] or '').strip() if len(row) > 15 else ''
    alc_dur = str(row[17] or '').strip() if len(row) > 17 else ''
    alc_lvl = str(row[18] or '').strip() if len(row) > 18 else ''
    alc_cost = str(row[19] or '').strip() if len(row) > 19 else ''
    if alc_name and alc_eff:
        alcohol.append({
            "name": alc_name,
            "effect": alc_eff,
            "duration": alc_dur,
            "level": alc_lvl,
            "cost": alc_cost
        })
        
    # Food (col 22..29)
    f_name = str(row[22] or '').strip() if len(row) > 22 else ''
    f_tier = str(row[23] or '').strip() if len(row) > 23 else ''
    f_eff = str(row[24] or '').strip() if len(row) > 24 else ''
    f_rec = str(row[25] or '').strip() if len(row) > 25 else ''
    f_dur = str(row[27] or '').strip() if len(row) > 27 else ''
    f_cost = str(row[29] or '').strip() if len(row) > 29 else ''
    if f_name and f_eff:
        food.append({
            "name": f_name,
            "tier": f_tier,
            "effect": f_eff,
            "recipe": f_rec,
            "duration": f_dur,
            "cost": f_cost
        })
        
    # Scrolls (col 31..41)
    s_name = str(row[31] or '').strip() if len(row) > 31 else ''
    s_eff = str(row[32] or '').strip() if len(row) > 32 else ''
    s_dur = str(row[33] or '').strip() if len(row) > 33 else ''
    s_cost = str(row[34] or '').strip() if len(row) > 34 else ''
    s_gems = str(row[35] or '').strip() if len(row) > 35 else ''
    if s_name and s_eff:
        scrolls.append({
            "name": s_name,
            "effect": s_eff,
            "duration": s_dur,
            "cost": s_cost,
            "gem_req": s_gems
        })
        
    # Staves (col 43..53)
    st_name = str(row[43] or '').strip() if len(row) > 43 else ''
    st_eff = str(row[44] or '').strip() if len(row) > 44 else ''
    st_mat = str(row[42] or '').strip() if len(row) > 42 else ''
    st_cost = str(row[47] or '').strip() if len(row) > 47 else ''
    if st_name and st_eff:
        staves.append({
            "name": st_name,
            "material": st_mat,
            "effect": st_eff,
            "cost": st_cost
        })

consumables_data = {
    "potions": potions,
    "poisons": poisons,
    "alcohol": alcohol,
    "food": food,
    "scrolls": scrolls,
    "staves": staves
}
with open(f"{DATA_DIR}/consumables.json", "w", encoding="utf-8") as f:
    json.dump(consumables_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 9. ENCHANTMENTS
# -------------------------------------------------------------
print("Exporting enchantments...")
ws_ench = wb['✨Зачарования']
ench_effects = []
current_type = "Экипировка"
for r, row in enumerate(ws_ench.iter_rows(values_only=True)):
    if r < 5:
        continue
    c2 = str(row[2] or '').strip() if len(row) > 2 else ''
    if c2 in ['Экипировка', 'Оружие']:
        current_type = c2
    name = str(row[3] or '').strip() if len(row) > 3 else ''
    if not name:
        continue
    t1 = str(row[7] or '').strip() if len(row) > 7 else ''
    t2 = str(row[8] or '').strip() if len(row) > 8 else ''
    t3 = str(row[9] or '').strip() if len(row) > 9 else ''
    t4 = str(row[10] or '').strip() if len(row) > 10 else ''
    t5 = str(row[11] or '').strip() if len(row) > 11 else ''
    desc = str(row[46] or '').strip() if len(row) > 46 else ''
    
    # Check slots/sources
    ench_effects.append({
        "type": current_type,
        "name": name,
        "tiers": [t1, t2, t3, t4, t5],
        "description": desc
    })

with open(f"{DATA_DIR}/enchantments.json", "w", encoding="utf-8") as f:
    json.dump(ench_effects, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 10. SPELLS & THU'UMS
# -------------------------------------------------------------
print("Exporting spells & shouts...")
ws_spells = wb['🌈ЗаклТуумы']
schools = {
    "Разрушение": [],
    "Колдовство": [],
    "Восстановление": [],
    "Изменение": [],
    "Иллюзия": [],
    "Вампиризм": [],
    "Туумы": []
}
current_school = "Разрушение"
current_tier = "Новичок"

for r, row in enumerate(ws_spells.iter_rows(values_only=True)):
    if r < 4:
        continue
    c11 = str(row[11] or '').strip() if len(row) > 11 else ''
    c5 = str(row[5] or '').strip() if len(row) > 5 else ''
    c2 = str(row[2] or '').strip() if len(row) > 2 else ''
    
    if "Колдовство" in c11:
        current_school = "Колдовство"
    elif "Восстановление" in c11:
        current_school = "Восстановление"
    elif "Изменение" in c11:
        current_school = "Изменение"
    elif "Иллюзия" in c11:
        current_school = "Иллюзия"
    elif "Заклинания вампира" in c11:
        current_school = "Вампиризм"
    elif "Туум" in c11 or "Крик" in c11 or r > 445:
        current_school = "Туумы"
        
    if c2 in ['Новичок', 'Ученик', 'Адепт', 'Эксперт', 'Мастер']:
        current_tier = c2
        continue
        
    name = c2
    if not name or name in ['Название', 'Сто-сть']:
        continue
        
    cost = str(row[3] or '').strip() if len(row) > 3 else ''
    dual = str(row[4] or '').strip() if len(row) > 4 else ''
    dmg = str(row[6] or '').strip() if len(row) > 6 else ''
    effect = str(row[11] or '').strip() if len(row) > 11 else ''
    form_id = str(row[18] or '').strip() if len(row) > 18 else ''
    cast_time = str(row[13] or '').strip() if len(row) > 13 else ''
    price = str(row[12] or '').strip() if len(row) > 12 else ''
    
    schools[current_school].append({
        "name": name,
        "school": current_school,
        "tier": current_tier,
        "cost": cost,
        "dual_cast": dual,
        "power": dmg,
        "effect": effect,
        "form_id": form_id,
        "cast_time": cast_time,
        "price": price
    })

with open(f"{DATA_DIR}/spells_shouts.json", "w", encoding="utf-8") as f:
    json.dump(schools, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 11. SUMMONS
# -------------------------------------------------------------
print("Exporting summons...")
ws_sum = wb['👻Самоны']
summons_list = []
current_sum_tier = "Новичок"

for r, row in enumerate(ws_sum.iter_rows(values_only=True)):
    if r < 4:
        continue
    c2 = str(row[2] or '').strip() if len(row) > 2 else ''
    if 'Новичок' in c2:
        current_sum_tier = "Новичок"
    elif 'Ученик' in c2:
        current_sum_tier = "Ученик"
    elif 'Адепт' in c2:
        current_sum_tier = "Адепт"
    elif 'Эксперт' in c2:
        current_sum_tier = "Эксперт"
    elif 'Мастер' in c2:
        current_sum_tier = "Мастер"
        
    if not c2 or c2.startswith('база') or 'Новичок' in c2 or 'Ученик' in c2 or 'Адепт' in c2 or 'Эксперт' in c2 or 'Мастер' in c2:
        continue
        
    cost = str(row[3] or '').strip() if len(row) > 3 else ''
    lvl = str(row[8] or '').strip() if len(row) > 8 else ''
    hp = str(row[9] or '').strip() if len(row) > 9 else ''
    stamina = str(row[10] or '').strip() if len(row) > 10 else ''
    magicka = str(row[11] or '').strip() if len(row) > 11 else ''
    dmg = str(row[12] or '').strip() if len(row) > 12 else ''
    armor = str(row[19] or '').strip() if len(row) > 19 else ''
    res_fire = str(row[20] or '').strip() if len(row) > 20 else ''
    res_frost = str(row[21] or '').strip() if len(row) > 21 else ''
    res_shock = str(row[22] or '').strip() if len(row) > 22 else ''
    res_chaos = str(row[23] or '').strip() if len(row) > 23 else ''
    abilities = str(row[27] or '').strip() if len(row) > 27 else ''
    
    summons_list.append({
        "name": c2,
        "tier": current_sum_tier,
        "cost": cost,
        "level": lvl,
        "health": hp,
        "stamina": stamina,
        "magicka": magicka,
        "damage": dmg,
        "armor": armor,
        "resistances": {
            "fire": res_fire,
            "frost": res_frost,
            "shock": res_shock,
            "chaos": res_chaos
        },
        "abilities": abilities
    })

with open(f"{DATA_DIR}/summons.json", "w", encoding="utf-8") as f:
    json.dump(summons_list, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 12. EQUIPMENT
# -------------------------------------------------------------
print("Exporting equipment...")
ws_eq = wb['⚔️Снаряжение']
weapons = []
armor_sets = []
cur_material = "Железное"

for r, row in enumerate(ws_eq.iter_rows(values_only=True)):
    if r < 3:
        continue
    c1 = str(row[1] or '').strip() if len(row) > 1 else ''
    if c1 and any(m in c1 for m in ['Железное', 'Стальное', 'Серебряное', 'Орочье', 'Гномье', 'Эльфийское', 'Стеклянное', 'Эбонитовое', 'Сталгримовое', 'Даэдрическое', 'Драконье']):
        cur_material = c1
        continue
    
    # Weapon row
    w_type = c1
    w_dmg = str(row[5] or '').strip() if len(row) > 5 else ''
    w_speed = str(row[6] or '').strip() if len(row) > 6 else ''
    w_dps = str(row[7] or '').strip() if len(row) > 7 else ''
    w_reach = str(row[8] or '').strip() if len(row) > 8 else ''
    w_weight = str(row[9] or '').strip() if len(row) > 9 else ''
    w_cost = str(row[11] or '').strip() if len(row) > 11 else ''
    
    if w_type and w_dmg and w_type != 'Тип' and w_type != '0.5333333333':
        weapons.append({
            "material": cur_material,
            "type": w_type,
            "damage": w_dmg,
            "speed": w_speed,
            "dps": w_dps,
            "reach": w_reach,
            "weight": w_weight,
            "cost": w_cost
        })

with open(f"{DATA_DIR}/equipment.json", "w", encoding="utf-8") as f:
    json.dump({"weapons": weapons}, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 13. UNIQUES
# -------------------------------------------------------------
print("Exporting uniques...")
ws_u = wb['💎Уники']
uniques_list = []
cur_cat = "Уникальные предметы"

for r, row in enumerate(ws_u.iter_rows(values_only=True)):
    if r < 2:
        continue
    c2 = str(row[2] or '').strip() if len(row) > 2 else ''
    c3 = str(row[3] or '').strip() if len(row) > 3 else ''
    
    if c2 and not c3:
        cur_cat = c2
        continue
        
    if not c2 or c2 == 'Название':
        continue
        
    mat = str(row[4] or '').strip() if len(row) > 4 else ''
    tier_ench = str(row[5] or '').strip() if len(row) > 5 else ''
    ench_count = str(row[6] or '').strip() if len(row) > 6 else ''
    loc_tier = str(row[7] or '').strip() if len(row) > 7 else ''
    stats = str(row[8] or '').strip() if len(row) > 8 else ''
    weight = str(row[9] or '').strip() if len(row) > 9 else ''
    price = str(row[10] or '').strip() if len(row) > 10 else ''
    location = str(row[13] or '').strip() if len(row) > 13 else ''
    form_id = str(row[14] or '').strip() if len(row) > 14 else ''
    
    uniques_list.append({
        "category": cur_cat,
        "name": c2,
        "effect": c3,
        "material": mat,
        "location_tier": loc_tier,
        "stats": stats,
        "weight": weight,
        "price": price,
        "location": location,
        "form_id": form_id
    })

with open(f"{DATA_DIR}/uniques.json", "w", encoding="utf-8") as f:
    json.dump(uniques_list, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 14. DIFFICULTY
# -------------------------------------------------------------
print("Exporting difficulty...")
diff_data = {
    "modes": [
        {"id": "adventure", "name": "Приключение", "badge": "assets/img/image34.jpg", "description": "Сбалансированный режим для спокойного изучения мира"},
        {"id": "tactics", "name": "Тактика", "badge": "assets/img/image33.jpg", "description": "Повышенные требования к позиционированию и ресурсам"},
        {"id": "heroic", "name": "Героический", "badge": "assets/img/image36.jpg", "description": "Серьезный вызов для опытных игроков"},
        {"id": "gods", "name": "Испытание богов", "badge": "assets/img/image35.jpg", "description": "Максимальный бескомпромиссный хардкор"}
    ],
    "multipliers": [
        {"param": "Наносимый игроком урон", "adventure": "2.0x", "tactics": "1.5x", "heroic": "1.2x", "gods": "1.0x"},
        {"param": "Получаемый игроком урон", "adventure": "0.25x", "tactics": "0.5x", "heroic": "0.8x", "gods": "1.0x"},
        {"param": "Получаемый опыт", "adventure": "3.0x", "tactics": "2.0x", "heroic": "1.5x", "gods": "1.0x"},
        {"param": "Скорость атак противников", "adventure": "-20%", "tactics": "-15%", "heroic": "-10%", "gods": "0%"},
        {"param": "Скорость полета снарядов противников", "adventure": "-20%", "tactics": "-10%", "heroic": "0%", "gods": "0%"},
        {"param": "Скорость разворота противников", "adventure": "-50%", "tactics": "-25%", "heroic": "0%", "gods": "0%"},
        {"param": "Регенерация здоровья противников", "adventure": "-75%", "tactics": "-50%", "heroic": "-25%", "gods": "0%"},
        {"param": "Длительность оглушения противников", "adventure": "+100%", "tactics": "+50%", "heroic": "+25%", "gods": "0%"},
        {"param": "Цены уроков, навыков и подношений", "adventure": "-50%", "tactics": "-25%", "heroic": "0%", "gods": "0%"}
    ],
    "rules": [
        {"rule": "Запрет смены камня-хранителя (кроме дерева)", "adventure": False, "tactics": False, "heroic": False, "gods": True},
        {"rule": "Запрет быстрых перемещений", "adventure": False, "tactics": False, "heroic": False, "gods": True},
        {"rule": "Ограничение на покупку очков навыков", "adventure": False, "tactics": False, "heroic": False, "gods": True},
        {"rule": "Увеличенный урон от падений", "adventure": False, "tactics": False, "heroic": False, "gods": True},
        {"rule": "Запрет сохранений в бою", "adventure": False, "tactics": False, "heroic": True, "gods": True},
        {"rule": "Запрет боя на лошади", "adventure": False, "tactics": True, "heroic": True, "gods": True},
        {"rule": "Драугры бесплотны при пробуждении", "adventure": False, "tactics": True, "heroic": True, "gods": True}
    ]
}
with open(f"{DATA_DIR}/difficulty.json", "w", encoding="utf-8") as f:
    json.dump(diff_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 15. BLACK BOOKS
# -------------------------------------------------------------
print("Exporting black books...")
black_books_data = [
    {
        "name": "Острое перо",
        "location": "Нчардак, в ходе выполнения задания «Путь знания»",
        "powers": [
            {"title": "Мощь Языков", "effect": "Урон ту'умов +10%"},
            {"title": "Техника медитации", "effect": "Сила ту'умов-аур +50%"},
            {"title": "Секреты тренировок", "effect": "Скорость восстановления ту'умов +20%"}
        ]
    },
    {
        "name": "Филамент и филигрань",
        "location": "Курган Колбьорн, по завершению квеста «Раскопки» после победы над Азидалом",
        "powers": [
            {"title": "Живучесть", "effect": "Регенерация здоровья +100%"},
            {"title": "Выдержка", "effect": "Регенерация запаса сил +100%"},
            {"title": "Отстраненность", "effect": "Регенерация магии +100%"}
        ]
    },
    {
        "name": "Скрытый сумрак",
        "location": "Тель-Митрин, в комнате с зачарователем посохов после квеста «Проблема с управителем»",
        "powers": [
            {"title": "Таинство ударов", "effect": "Расход атрибута на атаки и стрельбу -20%"},
            {"title": "Таинство чар", "effect": "Расход атрибута на заклинания -10%"},
            {"title": "Таинство стойки", "effect": "Отражение урона +10%"}
        ]
    },
    {
        "name": "Болезненный регент",
        "location": "Курган Белого Хребта",
        "powers": [
            {"title": "Забытые отвары", "effect": "Эффективность зелий усиления +20%"},
            {"title": "Древние снадобья", "effect": "Эффективность восполняющих зелий +4 ед/сек"},
            {"title": "Тонкий вкус", "effect": "Эффективность еды +60%"}
        ]
    },
    {
        "name": "Ветры перемен",
        "location": "Курган Бладскал (шахта Вороньей Скалы) после победы над Закрисошем",
        "powers": [
            {"title": "Нерушимость", "effect": "Здоровье +30 ед."},
            {"title": "Натренированность", "effect": "Запас сил +30 ед."},
            {"title": "Одаренность", "effect": "Магия +30 ед."}
        ]
    },
    {
        "name": "Нерассказанные легенды",
        "location": "Бенконгерик, за стеной Слов Силы",
        "powers": [
            {"title": "Стойкость", "effect": "Броня +100 ед."},
            {"title": "Закаленность", "effect": "Сопротивление стихиям +100 ед."},
            {"title": "Расторопность", "effect": "Скорость +10 ед."}
        ]
    },
    {
        "name": "Пробуждающий сон",
        "location": "Храм Мирака (Святилище Апокрифа)",
        "powers": [
            {"title": "Сброс способностей", "effect": "Позволяет заново перераспределить очки навыков в ветках"}
        ]
    }
]
with open(f"{DATA_DIR}/black_books.json", "w", encoding="utf-8") as f:
    json.dump(black_books_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 16. QUEST REWARDS
# -------------------------------------------------------------
print("Exporting quest rewards...")
ws_q = wb['💰Награды за квесты']
quests_list = []
for r, row in enumerate(ws_q.iter_rows(values_only=True)):
    if r < 31:
        continue
    name1 = str(row[3] or '').strip() if len(row) > 3 else ''
    chain = str(row[7] or '').strip() if len(row) > 7 else ''
    coins1 = str(row[9] or '').strip() if len(row) > 9 else ''
    if name1 and coins1 and name1 != 'Наименование':
        quests_list.append({
            "name": name1,
            "chain": chain if chain else "Сюжетные задания",
            "coins": coins1
        })
        
    name2 = str(row[13] or '').strip() if len(row) > 13 else ''
    coins2 = str(row[16] or '').strip() if len(row) > 16 else ''
    if name2 and coins2 and name2 != 'Наименование':
        quests_list.append({
            "name": name2,
            "chain": "Квесты Даэдра / Дополнительные",
            "coins": coins2
        })

quest_data = {
    "total_coins": 2150,
    "exchange_info": "Древние нордские монеты обмениваются на очки способностей на алтарях драконорожденного.",
    "quests": quests_list
}
with open(f"{DATA_DIR}/quest_rewards.json", "w", encoding="utf-8") as f:
    json.dump(quest_data, f, ensure_ascii=False, indent=2)

# -------------------------------------------------------------
# 17. META & GLOBAL SEARCH INDEX
# -------------------------------------------------------------
print("Generating search index...")
search_index = []

# Add base
for s in base_sections:
    search_index.append({
        "title": s['title'],
        "category": "База и механики",
        "snippet": s['content'][:140],
        "target": "base",
        "id": s['id']
    })

# Add stones
for st in races_stones_data["standing_stones"]:
    search_index.append({
        "title": f"Камень: {st['name']}",
        "category": "Расы и Камни",
        "snippet": st['description'][:140],
        "target": "races_stones"
    })

# Add aedra
for a in aedra_list:
    search_index.append({
        "title": f"Аэдра: {a['name']}",
        "category": "Аэдра",
        "snippet": f"Дар: {a['gift']} | Благословение: {a['blessing']}"[:140],
        "target": "aedra"
    })

# Add daedra
for d in daedra_list:
    search_index.append({
        "title": f"Даэдра: {d['name']}",
        "category": "Даэдра",
        "snippet": f"Дар: {d['gift']}"[:140],
        "target": "daedra"
    })

# Add perks
for tree in perks_data:
    for p in tree['perks']:
        search_index.append({
            "title": f"{p['name']} ({tree['name']})",
            "category": f"Перки: {tree['name']}",
            "snippet": f"Уровень: {p['level']} | {p['description']}"[:140],
            "target": "perks"
        })

# Add uniques
for u in uniques_list:
    search_index.append({
        "title": u['name'],
        "category": f"Уники: {u['category']}",
        "snippet": f"{u['effect']} (ID: {u['form_id']})"[:140],
        "target": "uniques"
    })

# Add spells
for sch, sp_list in schools.items():
    for sp in sp_list:
        search_index.append({
            "title": f"{sp['name']} ({sch})",
            "category": f"Заклинания: {sch}",
            "snippet": f"Тир: {sp['tier']} | {sp['effect']}"[:140],
            "target": "spells_shouts"
        })

with open(f"{DATA_DIR}/search_index.json", "w", encoding="utf-8") as f:
    json.dump(search_index, f, ensure_ascii=False)

print(f"All data successfully exported! Total search items: {len(search_index)}")
